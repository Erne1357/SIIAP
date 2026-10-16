"""DocumentGenerationService: variables, template selection and DOCX/HTML rendering."""

import sys
import types
from pathlib import Path

import pytest
from docx import Document
from freezegun import freeze_time

from app import db
from app.models.academic_period import AcademicPeriod
from app.models.document_template import DocumentTemplate
from app.services.document_generation_service import DocumentGenerationService as DG, _fecha_larga
from tests.support.factories import enroll, make_academic_period, make_program, make_user


@pytest.fixture
def env(app, tmp_path, monkeypatch):
    """Template store and instance path redirected into tmp_path."""
    store = tmp_path / 'templates_sys'
    store.mkdir()
    app.instance_path = str(tmp_path / 'instance')
    app.config['TEMPLATES_SYS_FOLDER'] = str(store)
    coord = make_user('program_admin', first_name='Ana', last_name='Coord')
    program = make_program(coord, name='Maestría en Pruebas')
    program.program_level = 'Maestría'
    student = make_user('student', first_name='Luis', last_name='Pérez', mother_last_name='Gómez')
    student.curp = 'PEGL950101HCHRMS09'
    student.control_number = 'A123'
    enroll(student, program, 'accepted')
    period = make_academic_period('20261')
    db.session.commit()
    return types.SimpleNamespace(store=store, tmp=tmp_path, program=program, student=student,
                                 period=period, coord=coord)


def _docx(path: Path, body: str, table: str | None = None, header: str | None = None,
          footer: str | None = None):
    doc = Document()
    doc.add_paragraph(body)
    if table:
        t = doc.add_table(rows=1, cols=1)
        t.rows[0].cells[0].paragraphs[0].add_run(table)
    if header:
        doc.sections[0].header.paragraphs[0].add_run(header)
    if footer:
        doc.sections[0].footer.paragraphs[0].add_run(footer)
    doc.save(str(path))


def _register(env, name, file_type='docx', *, program_id='self', doc_type='acceptance_letter',
              active=True):
    t = DocumentTemplate(
        program_id=env.program.id if program_id == 'self' else program_id,
        document_type=doc_type, name=name, file_path=name, file_type=file_type)
    t.is_active = active
    db.session.add(t)
    db.session.commit()
    return t


def _text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        parts += [c.text for r in t.rows for c in r.cells]
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs] + [p.text for p in s.footer.paragraphs]
    return '\n'.join(parts)


class TestVariables:
    def test_full_variable_set(self, env):
        v = DG.get_variables(env.student, env.program, env.period)
        assert v['student_name'] == 'Luis Pérez Gómez'
        assert v['student_email'] == env.student.email
        assert v['program_name'] == 'Maestría en Pruebas' and v['program_level'] == 'Maestría'
        assert v['period_code'] == '20261' and v['period_name'] == env.period.name
        assert v['coordinator_name'] == 'Ana Coord'
        assert v['student_curp'] == 'PEGL950101HCHRMS09' and v['control_number'] == 'A123'

    def test_every_documented_variable_exists(self):
        assert set(DG.VARIABLE_MAP) == {
            'student_name', 'student_curp', 'student_email', 'program_name', 'program_level',
            'period_code', 'period_name', 'acceptance_date', 'coordinator_name',
            'current_date', 'control_number'}

    def test_name_without_mother_last_name_has_no_trailing_space(self, env):
        env.student.mother_last_name = None
        assert DG.get_variables(env.student, env.program)['student_name'] == 'Luis Pérez'

    def test_defaults_without_period_level_or_numbers(self, env):
        env.program.program_level = None
        env.student.curp = None
        env.student.control_number = None
        v = DG.get_variables(env.student, env.program, None)
        assert v['program_level'] == 'Posgrado'
        assert v['period_code'] == '' and v['period_name'] == ''
        assert v['student_curp'] == '' and v['control_number'] == ''

    def test_dates_are_spanish_and_deterministic(self, env):
        with freeze_time('2026-07-31 12:00:00'):
            v = DG.get_variables(env.student, env.program)
        assert v['acceptance_date'] == '31 de julio de 2026'
        assert v['current_date'] == '31/07/2026'

    @pytest.mark.parametrize('month,name', [(1, 'enero'), (3, 'marzo'), (9, 'septiembre'), (12, 'diciembre')])
    def test_fecha_larga_months(self, month, name):
        from datetime import datetime
        assert _fecha_larga(datetime(2026, month, 5)) == f'5 de {name} de 2026'


class TestGenerateDocx:
    def test_substitutes_in_body_table_header_footer(self, env):
        _docx(env.store / 'carta.docx', 'Estimado {{student_name}}',
              table='Programa: {{program_name}}', header='{{period_code}}',
              footer='Atte. {{coordinator_name}}')
        _register(env, 'carta.docx')
        r = DG.generate(env.student.id, env.program.id, 'acceptance_letter', env.period.id)
        text = _text(r['output_path'])
        assert 'Estimado Luis Pérez Gómez' in text
        assert 'Programa: Maestría en Pruebas' in text
        assert '20261' in text and 'Atte. Ana Coord' in text
        assert '{{' not in text

    def test_result_shape(self, env):
        _docx(env.store / 'c.docx', 'x')
        _register(env, 'c.docx', doc_type='enrollment_confirmation')
        with freeze_time('2026-03-01 10:20:30'):
            r = DG.generate(env.student.id, env.program.id, 'enrollment_confirmation')
        assert r['file_type'] == 'docx'
        assert r['filename'] == 'enrollment-confirmation_20260301_042030.docx' or \
            r['filename'].startswith('enrollment-confirmation_20260301_')
        assert Path(r['output_path']).exists()
        assert str(env.student.id) in r['output_path']

    def test_unknown_placeholder_left_untouched(self, env):
        _docx(env.store / 'c.docx', 'Hola {{no_existe}} {{student_name}}')
        _register(env, 'c.docx')
        text = _text(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        assert '{{no_existe}}' in text and 'Luis' in text

    def test_repeated_placeholder_all_replaced(self, env):
        _docx(env.store / 'c.docx', '{{program_name}} / {{program_name}}')
        _register(env, 'c.docx')
        text = _text(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        assert text.count('Maestría en Pruebas') == 2

    def test_template_file_is_not_modified(self, env):
        _docx(env.store / 'c.docx', 'Hola {{student_name}}')
        before = (env.store / 'c.docx').read_bytes()
        _register(env, 'c.docx')
        DG.generate(env.student.id, env.program.id, 'acceptance_letter')
        assert (env.store / 'c.docx').read_bytes() == before

    def test_value_containing_braces_is_inserted_literally(self, env):
        env.student.first_name = '{{program_name}}'
        _docx(env.store / 'c.docx', '{{student_name}}')
        _register(env, 'c.docx')
        text = _text(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        assert 'Pérez' in text

    # Regression: placeholders split across Word runs were left unreplaced (run-by-run replacement).
    def test_placeholder_split_across_runs(self, env):
        doc = Document()
        p = doc.add_paragraph()
        first = p.add_run('Hola {{stud')
        first.bold = True
        p.add_run('ent_name}}')
        doc.save(str(env.store / 'c.docx'))
        _register(env, 'c.docx')
        out = Document(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        para = out.paragraphs[0]
        assert para.text == 'Hola Luis Pérez Gómez'
        assert para.runs[0].text == 'Hola Luis Pérez Gómez' and para.runs[0].bold is True

    def test_split_placeholder_in_nested_table_header_and_footer(self, env):
        doc = Document()
        doc.add_paragraph('cuerpo')
        inner = doc.add_table(rows=1, cols=1).rows[0].cells[0].add_table(rows=1, cols=1)
        cp = inner.rows[0].cells[0].paragraphs[0]
        cp.add_run('{{program_')
        cp.add_run('name}}')
        hp = doc.sections[0].header.paragraphs[0]
        hp.add_run('{{period')
        hp.add_run('_code}}')
        fp = doc.sections[0].footer.paragraphs[0]
        fp.add_run('Atte. {{')
        fp.add_run('coordinator_name}}')
        doc.save(str(env.store / 'c.docx'))
        _register(env, 'c.docx')
        r = DG.generate(env.student.id, env.program.id, 'acceptance_letter', env.period.id)
        out = Document(r['output_path'])
        nested = out.tables[0].rows[0].cells[0].tables[0].rows[0].cells[0]
        assert nested.text == 'Maestría en Pruebas'
        assert out.sections[0].header.paragraphs[0].text == '20261'
        assert out.sections[0].footer.paragraphs[0].text == 'Atte. Ana Coord'

    def test_runs_keep_their_formatting_when_not_split(self, env):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('{{student_name}}').bold = True
        p.add_run(' / ')
        p.add_run('{{program_name}}').italic = True
        doc.save(str(env.store / 'c.docx'))
        _register(env, 'c.docx')
        out = Document(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        runs = out.paragraphs[0].runs
        assert [r.text for r in runs] == ['Luis Pérez Gómez', ' / ', 'Maestría en Pruebas']
        assert runs[0].bold is True and runs[2].italic is True


class TestTemplateSelection:
    def test_program_template_beats_global(self, env):
        _docx(env.store / 'global.docx', 'GLOBAL')
        _docx(env.store / 'prog.docx', 'ESPECIFICA')
        _register(env, 'global.docx', program_id=None)
        _register(env, 'prog.docx')
        text = _text(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        assert 'ESPECIFICA' in text

    def test_global_used_when_no_specific(self, env):
        _docx(env.store / 'global.docx', 'GLOBAL')
        _register(env, 'global.docx', program_id=None)
        text = _text(DG.generate(env.student.id, env.program.id, 'acceptance_letter')['output_path'])
        assert 'GLOBAL' in text

    def test_inactive_template_ignored(self, env):
        _docx(env.store / 'prog.docx', 'X')
        _register(env, 'prog.docx', active=False)
        with pytest.raises(ValueError, match='No hay plantilla activa'):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')

    def test_other_programs_template_not_used(self, env):
        other = make_program()
        db.session.commit()
        _docx(env.store / 'o.docx', 'X')
        t = DocumentTemplate(program_id=other.id, document_type='acceptance_letter', name='o',
                             file_path='o.docx', file_type='docx')
        db.session.add(t)
        db.session.commit()
        with pytest.raises(ValueError):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')

    def test_other_document_type_not_used(self, env):
        _docx(env.store / 'p.docx', 'X')
        _register(env, 'p.docx', doc_type='course_schedule')
        with pytest.raises(ValueError):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')

    def test_missing_file_on_disk(self, env):
        _register(env, 'ghost.docx')
        with pytest.raises(ValueError, match='no existe en disco'):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')

    def test_unsupported_file_type(self, env):
        (env.store / 'x.txt').write_text('x')
        _register(env, 'x.txt', file_type='txt')
        with pytest.raises(ValueError, match='no soportado'):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')

    def test_unknown_student_404(self, env):
        from werkzeug.exceptions import NotFound
        with pytest.raises(NotFound):
            DG.generate(99999, env.program.id, 'acceptance_letter')

    def test_unknown_program_404(self, env):
        from werkzeug.exceptions import NotFound
        with pytest.raises(NotFound):
            DG.generate(env.student.id, 99999, 'acceptance_letter')


class TestGenerateHtml:
    @pytest.fixture
    def fake_weasy(self, monkeypatch):
        captured = {}

        class FakeHTML:
            def __init__(self, string, base_url=None):
                captured['html'] = string
                captured['base_url'] = base_url

            def write_pdf(self, path):
                Path(path).write_bytes(b'%PDF-fake')

        mod = types.ModuleType('weasyprint')
        mod.HTML = FakeHTML
        mod.CSS = object
        monkeypatch.setitem(sys.modules, 'weasyprint', mod)
        return captured

    def test_html_template_rendered_to_pdf(self, env, fake_weasy):
        (env.store / 'carta.html').write_text(
            '<h1>{{program_name}}</h1><p>{{student_name}} - {{current_date}}</p>', encoding='utf-8')
        _register(env, 'carta.html', file_type='html')
        with freeze_time('2026-03-01 18:00:00'):
            r = DG.generate(env.student.id, env.program.id, 'acceptance_letter')
        assert r['file_type'] == 'pdf' and r['filename'].endswith('.pdf')
        assert Path(r['output_path']).read_bytes().startswith(b'%PDF')
        assert '<h1>Maestría en Pruebas</h1>' in fake_weasy['html']
        assert 'Luis Pérez Gómez - 01/03/2026' in fake_weasy['html']
        assert '{{' not in fake_weasy['html']

    def test_base_url_is_template_directory(self, env, fake_weasy):
        (env.store / 'c.html').write_text('x', encoding='utf-8')
        _register(env, 'c.html', file_type='html')
        DG.generate(env.student.id, env.program.id, 'acceptance_letter')
        assert Path(fake_weasy['base_url']) == env.store

    def test_missing_weasyprint_gives_clear_error(self, env, monkeypatch):
        monkeypatch.setitem(sys.modules, 'weasyprint', None)
        (env.store / 'c.html').write_text('x', encoding='utf-8')
        _register(env, 'c.html', file_type='html')
        with pytest.raises(RuntimeError, match='weasyprint'):
            DG.generate(env.student.id, env.program.id, 'acceptance_letter')
